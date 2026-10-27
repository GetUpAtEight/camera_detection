"""把 docs/SOLUTION.md 转成一份可直接汇报的 Word 文档（docs/SOLUTION.docx）。

用法::

    python deploy/build_solution_docx.py

只用标准库 + python-docx。刻意保持"轻量 Markdown → DOCX"的直译：
标题层级、段落、无序/有序列表、表格、代码块、引用块、分隔线。
不追求完整 Markdown 语义，够读、够专业即可。
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "SOLUTION.md")
OUT = os.path.join(ROOT, "docs", "SOLUTION.docx")

BODY_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)


# --------------------------------------------------------------------------- 基础样式
def set_east_asian(run, font_name: str) -> None:
    """中文必须显式指定 w:eastAsia，否则 Word 会回退到默认字体。"""
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def shade(paragraph_or_cell, hex_color: str) -> None:
    el = paragraph_or_cell._element
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, color in (("Heading 1", 16, ACCENT), ("Heading 2", 13, ACCENT),
                              ("Heading 3", 11.5, RGBColor(0x2E, 0x5C, 0x8A))):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.font.name = BODY_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    # 页脚：说明文字 + 页码（顺序必须是"文字在前、页码在后"）
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = para.add_run("摄像头篡改检测系统 · 技术解决方案   ·   第 ")
    set_east_asian(r1, BODY_FONT)
    r1.font.size = Pt(8)
    r1.font.color.rgb = MUTED

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    inner = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    inner.append(rpr)
    txt = OxmlElement("w:t")
    txt.text = "1"
    inner.append(txt)
    fld.append(inner)
    para._p.append(fld)

    r2 = para.add_run(" 页")
    set_east_asian(r2, BODY_FONT)
    r2.font.size = Pt(8)
    r2.font.color.rgb = MUTED


# --------------------------------------------------------------------------- 行内解析
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]*\]\([^)]*\))")


def add_inline(paragraph, text: str) -> None:
    """把 **粗体** / `代码` / [文字](链接) 拆成 run，保留基本强调。"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_east_asian(run, BODY_FONT)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_east_asian(run, BODY_FONT)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.size = Pt(9.5)
            set_east_asian(run, MONO_FONT)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            label = token[1:token.index("]")]
            url = token[token.index("(") + 1:-1]
            run = paragraph.add_run(label if url.startswith("#") else f"{label} ({url})")
            set_east_asian(run, BODY_FONT)
            run.font.color.rgb = ACCENT
            if not url.startswith("#"):
                run.underline = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_east_asian(run, BODY_FONT)


def plain(text: str) -> str:
    """表格单元格里只要纯文本。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    return text.strip()


# --------------------------------------------------------------------------- 块级解析
def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    # 跨页时重复表头，长表格才不会"续页后不知道每列是什么"
    header_tr = table.rows[0]._tr
    tr_pr = header_tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.1
            text = row[j] if j < len(row) else ""
            run = para.add_run(plain(text))
            set_east_asian(run, BODY_FONT)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                shade(cell, "EAF1F8")
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]) -> None:
    if not lines:
        return
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.left_indent = Cm(0.4)
    para.paragraph_format.line_spacing = 1.0
    shade(para, "F5F5F5")
    for idx, line in enumerate(lines):
        run = para.add_run(line)
        set_east_asian(run, MONO_FONT)
        run.font.size = Pt(8.5)
        if idx < len(lines) - 1:
            run.add_break()


def add_quote(doc: Document, lines: list[str]) -> None:
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.6)
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run("│ ")
        set_east_asian(run, MONO_FONT)
        run.font.color.rgb = ACCENT
        add_inline(para, line)
        for r in para.runs[1:]:
            r.font.color.rgb = MUTED


def convert() -> None:
    with open(SRC, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    setup_document(doc)

    # 封面标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    run = title.add_run("摄像头篡改检测系统")
    set_east_asian(run, BODY_FONT)
    run.font.size = Pt(26)
    run.bold = True
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("技术解决方案")
    set_east_asian(run, BODY_FONT)
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(30)
    run = meta.add_run("v1.0 · 2026-10 · 面向客户交付与评审")
    set_east_asian(run, BODY_FONT)
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED

    doc.add_page_break()

    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    pending_list_para = None  # 用于把"缩进续行"接回上一条列表项

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line.rstrip())
            i += 1
            continue

        # 表格（注意：这里**不能** flush，否则每一行都会变成一张独立的表）
        if stripped.startswith("|"):
            cells = [c for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip()):
                i += 1
                continue
            table_rows.append([c.strip() for c in cells])
            i += 1
            continue
        flush_table()

        if not stripped:
            pending_list_para = None
            i += 1
            continue

        # 列表项的缩进续行：接回上一条，而不是另起一段
        if (pending_list_para is not None and re.match(r"^\s{2,}\S", line)
                and not re.match(r"^\s*[-*]\s", line) and not re.match(r"^\s*\d+\.\s", line)):
            add_inline(pending_list_para, stripped)
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            para = doc.add_paragraph()
            pPr = para._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "DDDDDD")
            border.append(bottom)
            pPr.append(border)
            i += 1
            continue

        # 标题（一级标题已在封面呈现，正文里跳过，避免重复大标题）
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = plain(m.group(2))
            if level == 1:
                i += 1
                continue
            doc.add_heading(text, level=min(level - 1, 3))
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_quote(doc, quote_lines)
            continue

        # 无序列表（手写符号 + 悬挂缩进，避免 Word 列表样式在长文档里跑偏）
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.85)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(2)
            run = para.add_run("•  ")
            set_east_asian(run, BODY_FONT)
            add_inline(para, m.group(1))
            pending_list_para = para
            i += 1
            continue

        # 有序列表（保留 Markdown 里的原始编号，避免跨列表累加）
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.95)
            pf.first_line_indent = Cm(-0.55)
            pf.space_after = Pt(2)
            run = para.add_run(f"{m.group(1)}.  ")
            set_east_asian(run, BODY_FONT)
            add_inline(para, m.group(2))
            pending_list_para = para
            i += 1
            continue

        # 普通段落
        pending_list_para = None
        para = doc.add_paragraph()
        add_inline(para, stripped)
        i += 1

    flush_table()
    if code_lines:
        add_code_block(doc, code_lines)

    doc.save(OUT)
    print(f"saved {OUT}")


def verify() -> int:
    doc = Document(OUT)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    tables = len(doc.tables)
    chars = sum(len(p.text) for p in doc.paragraphs)
    print(f"paragraphs={len(doc.paragraphs)} tables={tables} "
          f"headings={len(headings)} chars≈{chars}")
    missing = [t for t in ("1. 项目概述", "6. 算法设计：三层信号模型", "14. 测试与验证",
                           "17. 运维手册", "附录 A：关键参数表") if t not in headings]
    if missing:
        print("MISSING headings:", missing)
        return 1
    if tables < 20:
        print(f"WARN: only {tables} tables (expected >=20)")
        return 1
    print("verify OK")
    return 0


if __name__ == "__main__":
    convert()
    sys.exit(verify())
